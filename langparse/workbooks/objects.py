"""Source-grounded drawing facts and semantic blocks; no image inference or I/O."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from langparse.types import ParseDiagnostics
from langparse.workbooks.types import (
    SourceRef,
    WorkbookBlock,
    WorkbookIR,
    WorkbookSnapshot,
    stable_id,
)


@dataclass
class ChartSeries:
    title: str | None = None
    title_reference: str | None = None
    categories: str | None = None
    values: str | None = None
    x_values: str | None = None
    y_values: str | None = None
    source_refs: list[SourceRef] = field(default_factory=list)


@dataclass
class DrawingFact:
    object_id: str
    kind: str
    anchor: str | None
    title: str | None = None
    chart_types: list[str] = field(default_factory=list)
    series: list[ChartSeries] = field(default_factory=list)
    axes: list[dict[str, Any]] = field(default_factory=list)
    width: float | None = None
    height: float | None = None
    media_id: str | None = None
    enhancement_status: str = "not_requested"
    diagnostics: list[str] = field(default_factory=list)


def _text(value: Any) -> str | None:
    if isinstance(value, str):
        return value
    # DrawingML rich text can split a title over multiple runs/paragraphs.
    rich = getattr(getattr(value, "tx", None), "rich", None)
    paragraphs = getattr(rich, "p", [])
    lines = []
    for paragraph in paragraphs:
        runs = [getattr(run, "t", "") for run in getattr(paragraph, "r", [])]
        lines.append("".join(runs))
    return "\n".join(lines).strip() or None


def _formula(value: Any) -> str | None:
    for field_name in ("numRef", "strRef", "multiLvlStrRef"):
        reference = getattr(value, field_name, None)
        if reference is not None:
            return getattr(reference, "f", None)
    return None


def _source(formula: str, diagnostics: list[str], sheet_names: dict[str, str]) -> SourceRef | None:
    # External, union and dynamic expressions remain explicit raw facts.
    match = re.fullmatch(
        r"(?:'((?:[^']|'')+)'|([^'!]+))!(\$?[A-Z]+\$?[1-9][0-9]*(?::\$?[A-Z]+\$?[1-9][0-9]*)?)",
        formula,
    )
    if match and "[" not in formula:
        from langparse.workbooks.lineage import _bounds

        name = (match[1] or match[2]).replace("''", "'")
        cell_range = match[3].replace("$", "")
        try:
            _bounds(cell_range)
        except ValueError:
            pass
        else:
            if name.casefold() in sheet_names:
                return SourceRef(sheet_names[name.casefold()], cell_range)
    diagnostics.append(f"unresolved_chart_reference: {formula}")
    return None


def extract_objects(sheet: Any) -> list[DrawingFact]:
    from openpyxl.utils import get_column_letter

    sheet_names = {name.casefold(): name for name in sheet.parent.sheetnames}
    result = []
    for kind, objects in (("chart", sheet._charts), ("image", sheet._images)):
        for index, obj in enumerate(objects):
            anchor = obj.anchor if isinstance(obj.anchor, str) else None
            marker = getattr(obj.anchor, "_from", None)
            if marker is not None:
                anchor = f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
            fact = DrawingFact(stable_id("object", sheet.title, kind, str(index)), kind, anchor)
            if anchor is None:
                fact.diagnostics.append("unresolved_anchor")
            if kind == "image":
                fact.width, fact.height = obj.width, obj.height
                fact.media_id = stable_id("media", sheet.title, str(index))
            else:
                fact.title = _text(obj.title)
                for chart in getattr(obj, "_charts", [obj]):
                    fact.chart_types.append(type(chart).__name__)
                    for axis_name in ("x_axis", "y_axis", "z_axis"):
                        axis = getattr(chart, axis_name, None)
                        if axis is not None:
                            fact.axes.append(
                                {
                                    "axis": axis_name,
                                    "id": axis.axId,
                                    "title": _text(axis.title),
                                    "position": axis.axPos,
                                }
                            )
                    for series in chart.series:
                        tx = getattr(series, "tx", None)
                        item = ChartSeries(
                            title=getattr(tx, "v", None),
                            title_reference=_formula(tx),
                            categories=_formula(getattr(series, "cat", None)),
                            values=_formula(getattr(series, "val", None)),
                            x_values=_formula(getattr(series, "xVal", None)),
                            y_values=_formula(getattr(series, "yVal", None)),
                        )
                        for expression in (
                            item.title_reference,
                            item.categories,
                            item.values,
                            item.x_values,
                            item.y_values,
                        ):
                            if expression:
                                ref = _source(expression, fact.diagnostics, sheet_names)
                                if ref is not None:
                                    item.source_refs.append(ref)
                        fact.series.append(item)
            result.append(fact)
    return result


def attach_objects(
    snapshot: WorkbookSnapshot, ir: WorkbookIR, diagnostics: ParseDiagnostics
) -> None:
    """Add semantic objects after cell interpretation without claiming cell coverage."""
    total = covered = 0
    sheets = {sheet.name: sheet for sheet in ir.sheets}
    for source in snapshot.sheets:
        target = sheets[source.name]
        for obj in source.objects:
            total += 1
            anchor = obj.get("anchor")
            notes = list(obj.get("diagnostics", []))
            try:
                validate_object_source(snapshot, f"{source.name}!{anchor or ''}")
            except (TypeError, ValueError):
                notes.append("unresolved_anchor")
                diagnostics.warnings.append(f"Object on {source.name}: unresolved_anchor")
                diagnostics.status = "partial"
                continue
            refs = [SourceRef(source.name, anchor)]
            for series in obj.get("series", []):
                for raw in series.get("source_refs", []):
                    if raw["sheet_name"] in sheets:
                        ref = SourceRef(**raw)
                        if ref not in refs:
                            refs.append(ref)
                    else:
                        notes.append(f"missing_chart_sheet: {raw['sheet_name']}")
            if notes:
                diagnostics.status = "partial"
                diagnostics.warnings.extend(f"Object on {source.name}: {note}" for note in notes)
            target.blocks.append(
                WorkbookBlock(
                    block_id=obj.get("object_id")
                    or stable_id("object", source.name, anchor, str(total)),
                    kind=obj["kind"],
                    source_refs=refs,
                    metadata={"anchor": anchor, "object": obj},
                    diagnostics=[{"code": note} for note in notes],
                )
            )
            covered += 1
            kind = obj["kind"]
            diagnostics.block_count_by_kind[kind] = diagnostics.block_count_by_kind.get(kind, 0) + 1
    diagnostics.object_coverage_ratio = covered / total if total else 1.0
    if snapshot.metadata.get("drawing_read_errors"):
        diagnostics.status = "partial"
        diagnostics.object_coverage_ratio = None
        diagnostics.unsupported_features.extend(snapshot.metadata["drawing_read_errors"])


def render_object(block: WorkbookBlock) -> str:
    obj = block.metadata["object"]
    source = ", ".join(ref.key for ref in block.source_refs)
    title = obj.get("title") or obj["kind"].capitalize()
    lines = [f"<!-- source_ranges: {source} -->", f"**{title}** ({obj['kind']})"]
    for series in obj.get("series", []):
        references = [
            series[key]
            for key in ("categories", "values", "x_values", "y_values")
            if series.get(key)
        ]
        lines.append("- " + (series.get("title") or "Series") + ": " + ", ".join(references))
    if obj["kind"] == "image":
        lines.append(
            f"Image {obj.get('media_id')}; {obj.get('width')} × {obj.get('height')}; OCR/VLM not requested."
        )
    return "\n\n".join(lines)


def validate_object_source(snapshot: WorkbookSnapshot | None, source: str) -> None:
    """Drawing anchors can be outside used cells, but must be valid Excel coordinates."""
    from openpyxl.utils.cell import range_boundaries

    sheet_name, cell_range = source.rsplit("!", 1)
    if snapshot is None or not any(sheet.name == sheet_name for sheet in snapshot.sheets):
        raise ValueError(f"Object source references unknown sheet: {sheet_name}")
    try:
        left, top, right, bottom = range_boundaries(cell_range)
        if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
            raise ValueError("out of bounds")
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Invalid object source range: {source}") from exc
