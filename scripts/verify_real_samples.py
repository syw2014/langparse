"""Local, source-backed regression checks; never a production certification.

Run from the checkout with ``python -m scripts.verify_real_samples --help``.
Private manifests and reports belong in git-ignored directories. Source documents
are opened read-only. No cloud engine or workbook model is enabled by this runner.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import html
import json
import os
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from dataclasses import asdict, fields, is_dataclass
from importlib.metadata import PackageNotFoundError, version
from itertools import zip_longest
from pathlib import Path

import langparse
from langparse import __version__
from langparse.services.parse_service import ParseService
from langparse.workbooks.modeling import WorkbookDisambiguation


def _digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _json_default(value):
    # Do not deep-copy a large snapshot and both chunk profiles with asdict.
    if is_dataclass(value):
        return {field.name: getattr(value, field.name) for field in fields(value)}
    return str(value)


def _write_json(path, value):
    path = Path(path)
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "wt", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, default=_json_default)
        stream.write("\n")


def _unique_fields(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("Manifest contains duplicate fields")
        result[key] = value
    return result


def create_manifest(paths, destination):
    """Pin inputs only. Initialization never blesses parser-derived expectations."""
    destination = Path(destination).resolve()
    samples = []
    for index, path in enumerate(paths, 1):
        path = Path(path).resolve()
        if path.suffix.lower() not in {".xlsx", ".xlsm", ".pdf", ".csv"}:
            raise ValueError("Only OOXML workbooks, CSV and PDFs are supported")
        samples.append(
            {
                "id": f"sample-{index:03}",
                "path": os.path.relpath(path, destination.parent),
                "sha256": _digest(path),
                "semantic_checks": [],
            }
        )
    if not samples or len({s["path"] for s in samples}) != len(samples):
        raise ValueError("Provide nonempty, distinct sample paths")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("x", encoding="utf-8") as stream:
        json.dump(
            {"schema_version": 1, "split": "regression", "samples": samples},
            stream,
            ensure_ascii=False,
            indent=2,
        )


def _load_manifest(path):
    data = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=_unique_fields)
    if not isinstance(data, dict) or set(data) != {"schema_version", "split", "samples"}:
        raise ValueError("Invalid manifest fields")
    if (
        type(data["schema_version"]) is not int
        or data["schema_version"] != 1
        or data["split"] != "regression"
    ):
        raise ValueError("Expected schema_version=1 and split=regression")
    if not isinstance(data["samples"], list) or not data["samples"]:
        raise ValueError("Manifest must contain samples")
    ids, paths = set(), set()
    for sample in data["samples"]:
        if not isinstance(sample, dict) or set(sample) != {
            "id",
            "path",
            "sha256",
            "semantic_checks",
        }:
            raise ValueError("Invalid sample fields")
        if (
            not isinstance(sample["id"], str)
            or not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", sample["id"])
            or sample["id"] in ids
        ):
            raise ValueError("Sample IDs must be safe and unique")
        if not isinstance(sample["path"], str) or not sample["path"]:
            raise ValueError("Sample path must be a nonempty string")
        source = (path.parent / sample["path"]).resolve()
        if source in paths or source.suffix.lower() not in {".xlsx", ".xlsm", ".pdf", ".csv"}:
            raise ValueError("Unsupported or duplicate source")
        if not isinstance(sample["sha256"], str) or not re.fullmatch(
            r"[0-9a-f]{64}", sample["sha256"]
        ):
            raise ValueError("Expected a pinned SHA-256 digest")
        if not isinstance(sample["semantic_checks"], list):
            raise ValueError("semantic_checks must be a list")
        for check in sample["semantic_checks"]:
            _validate_check(check, source.suffix.lower())
        if source.suffix.lower() == ".csv" and len(sample["semantic_checks"]) != 1:
            raise ValueError("CSV requires exactly one explicit csv_table annotation")
        ids.add(sample["id"])
        paths.add(source)
    return data


def _validate_check(check, suffix):
    if not isinstance(check, dict):
        raise ValueError("Each semantic check must be an object")
    kind = check.get("kind")
    fields = {
        "header": {"kind", "sheet", "column", "source_cells"},
        "header_path": {"kind", "sheet", "column", "header_row", "expected"},
        "drawing_count": {"kind", "sheet", "expected"},
        "csv_table": {"kind", "delimiter", "encoding", "row_count", "column_count", "headers"},
        "row_role": {"kind", "sheet", "source_range", "expected"},
        "pdf_text": {"kind", "page", "text"},
        "pdf_table": {"kind", "page", "minimum"},
    }
    if kind not in fields or set(check) != fields[kind]:
        raise ValueError("Invalid semantic check fields")
    if kind == "csv_table":
        if (
            suffix != ".csv"
            or check["delimiter"] not in {",", "\t", ";", "|"}
            or check["encoding"] not in {"utf-8", "utf-8-sig", "gb18030", "utf-16"}
            or any(
                type(check[key]) is not int or check[key] < 1
                for key in ("row_count", "column_count")
            )
            or not isinstance(check["headers"], list)
            or len(check["headers"]) != check["column_count"]
            or any(not isinstance(value, str) for value in check["headers"])
        ):
            raise ValueError("Invalid CSV source annotation")
    elif kind in {"pdf_text", "pdf_table"}:
        if suffix != ".pdf" or type(check["page"]) is not int or check["page"] < 1:
            raise ValueError("Invalid PDF page check")
        if kind == "pdf_text" and (not isinstance(check["text"], str) or not check["text"].strip()):
            raise ValueError("Invalid PDF text check")
        if kind == "pdf_table" and (type(check["minimum"]) is not int or check["minimum"] < 1):
            raise ValueError("Invalid PDF table check")
    else:
        if suffix in {".pdf", ".csv"} or not isinstance(check["sheet"], str) or not check["sheet"]:
            raise ValueError("Invalid workbook check")
        if kind == "drawing_count":
            if type(check["expected"]) is not int or check["expected"] < 0:
                raise ValueError("Invalid drawing count")
            return
        if kind in {"header", "header_path"}:
            if not isinstance(check["column"], str) or not re.fullmatch("[A-Z]+", check["column"]):
                raise ValueError("Invalid header column")
            if kind == "header_path":
                if (
                    type(check["header_row"]) is not int
                    or check["header_row"] < 1
                    or not isinstance(check["expected"], list)
                    or any(not isinstance(v, str) or not v.strip() for v in check["expected"])
                ):
                    raise ValueError("Invalid header path")
                return
            if (
                not isinstance(check["source_cells"], list)
                or not check["source_cells"]
                or any(
                    not isinstance(c, str) or not re.fullmatch(r"[A-Z]+[1-9][0-9]*", c)
                    for c in check["source_cells"]
                )
            ):
                raise ValueError("Invalid header source cells")
        elif (
            not isinstance(check["source_range"], str)
            or not re.fullmatch(r"[A-Z]+[1-9][0-9]*:[A-Z]+[1-9][0-9]*", check["source_range"])
            or check["expected"]
            not in {
                "title",
                "context",
                "header",
                "data",
                "total",
                "section_header",
                "unknown",
                "repeated_title",
                "repeated_context",
                "repeated_header",
            }
        ):
            raise ValueError("Invalid row-role check")


def compare_workbook(path, parsed):
    """Read source facts directly; do not use LangParse's adapter or diagnostics."""
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    differences, nonempty, formulas, merges = [], 0, 0, 0
    snapshots = parsed.structure.snapshot.sheets
    try:
        if book.sheetnames != [s.name for s in snapshots]:
            differences.append({"code": "sheet_order_mismatch"})
        actual_sheets = {s.name: s for s in snapshots}
        for sheet in book:
            actual = actual_sheets.get(sheet.title)
            if actual is None:
                differences.append({"code": "missing_sheet", "sheet": sheet.title})
                continue
            for name, expected in {
                "visibility": sheet.sheet_state,
                "merged_ranges": sorted(str(r) for r in sheet.merged_cells.ranges),
                "hidden_rows": sorted(i for i, d in sheet.row_dimensions.items() if d.hidden),
                "hidden_columns": sorted(i for i, d in sheet.column_dimensions.items() if d.hidden),
            }.items():
                if getattr(actual, name) != expected:
                    differences.append({"code": f"{name}_mismatch", "sheet": sheet.title})
            merges += len(sheet.merged_cells.ranges)
            source_coords = set()
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    source_coords.add(cell.coordinate)
                    nonempty += 1
                    formula = cell.value if cell.data_type == "f" else None
                    formulas += formula is not None
                    target = actual.cells.get(cell.coordinate)
                    where = {"sheet": sheet.title, "cell": cell.coordinate}
                    if target is None:
                        differences.append({"code": "missing_cell", **where})
                        continue
                    for name, expected, code in (
                        ("raw_value", cell.value, "cell_value_mismatch"),
                        ("formula", formula, "formula_mismatch"),
                        (
                            "cached_value",
                            cached[sheet.title][cell.coordinate].value,
                            "cached_value_mismatch",
                        ),
                        ("number_format", cell.number_format, "number_format_mismatch"),
                    ):
                        if getattr(target, name) != expected:
                            differences.append({"code": code, **where})
            for coord, cell in actual.cells.items():
                if cell.raw_value is not None and coord not in source_coords:
                    differences.append(
                        {"code": "unexpected_cell", "sheet": sheet.title, "cell": coord}
                    )
    finally:
        book.close()
        cached.close()
    return {
        "status": "failed" if differences else "passed",
        "sheet_count": len(snapshots),
        "nonempty_cells": nonempty,
        "formula_cells": formulas,
        "merged_ranges": merges,
        "differences": differences,
    }


def check_chunks(parsed, chunks):
    rows_by_id = {
        row.row_id: row
        for sheet in parsed.structure.sheets
        for block in sheet.blocks
        if block.logical_table
        for row in block.logical_table.rows
        if row.role in {"data", "total"}
    }
    expected = list(rows_by_id)
    actual = [
        row_id
        for chunk in chunks
        if chunk.metadata.get("chunk_type") == "table_rows"
        for row_id in chunk.metadata.get("row_ids", [])
    ]
    missing, added = len(set(expected) - set(actual)), len(set(actual) - set(expected))
    duplicates = len(actual) - len(set(actual))
    without_refs = sum(not c.metadata.get("source_ranges") for c in chunks)
    payload_mismatches = 0
    for chunk in chunks:
        if chunk.metadata.get("chunk_type") != "table_rows":
            continue
        row_ids = chunk.metadata.get("row_ids", [])
        if any(row_id not in rows_by_id for row_id in row_ids):
            payload_mismatches += 1
            continue
        rows = [rows_by_id[row_id] for row_id in row_ids]
        payload_mismatches += chunk.structured_payload.get("rows") != [r.values for r in rows]
        payload_mismatches += chunk.structured_payload.get("roles") != [r.role for r in rows]
        payload_mismatches += chunk.metadata.get("source_ranges") != [
            r.source_ref.key for r in rows
        ]
        if chunk.metadata.get("chunk_profile") == "analysis":
            records = chunk.structured_payload.get("records", [])
            payload_mismatches += len(records) != len(rows)
            for record, row in zip(records, rows, strict=False):
                payload_mismatches += (
                    record.get("row_id") != row.row_id
                    or record.get("values") != row.values
                    or record.get("role") != row.role
                    or record.get("source_refs") != [row.source_ref.key]
                )
    return {
        "status": "failed"
        if missing or added or duplicates or without_refs or payload_mismatches or not chunks
        else "passed",
        "expected_rows": len(expected),
        "actual_rows": len(actual),
        "missing": missing,
        "unexpected": added,
        "duplicates": duplicates,
        "chunks_without_refs": without_refs,
        "payload_mismatches": payload_mismatches,
        "chunk_count": len(chunks),
    }


def _semantic_result(check, expected, actual):
    return {
        "check": check,
        "expected": expected,
        "actual": actual,
        "status": "passed" if expected == actual else "failed",
    }


def compare_csv(path, parsed, annotation):
    """Use the explicitly reviewed delimiter, independent of parser sniffing."""
    with Path(path).open(encoding=annotation["encoding"], newline="") as stream:
        expected = list(csv.reader(stream, delimiter=annotation["delimiter"], strict=True))
    if not expected or any(len(row) != annotation["column_count"] for row in expected):
        raise ValueError("Source CSV does not match the annotated rectangular schema")
    actual = [row for page in parsed.pages for table in page.tables for row in table["rows"]]
    differences = []
    for row_number, (left, right) in enumerate(zip_longest(expected, actual, fillvalue=[]), 1):
        for column, (wanted, got) in enumerate(zip_longest(left, right), 1):
            if wanted != got:
                differences.append(
                    {"row": row_number, "column": column, "expected": wanted, "actual": got}
                )
    semantic = [
        _semantic_result(
            {"kind": "csv_row_count", "location": Path(path).name},
            annotation["row_count"],
            len(actual),
        ),
        _semantic_result(
            {"kind": "csv_headers", "location": Path(path).name},
            annotation["headers"],
            actual[0] if actual else [],
        ),
        _semantic_result(
            {"kind": "csv_exact_cells", "location": Path(path).name}, 0, len(differences)
        ),
    ]
    return {
        "status": "failed"
        if differences or any(c["status"] == "failed" for c in semantic)
        else "passed",
        "source_rows": len(expected),
        "source_columns": annotation["column_count"],
        "differences": differences,
        "semantic_checks": semantic,
    }


def check_semantics(path, parsed, checks):
    """Only evaluate explicit, source-reviewed annotations. Never derive truth from output."""
    import openpyxl
    from openpyxl.utils import range_boundaries

    if Path(path).suffix.lower() == ".csv":
        return compare_csv(path, parsed, checks[0])["semantic_checks"]
    book = (
        openpyxl.load_workbook(path, data_only=False)
        if any(c["kind"] == "header" for c in checks)
        else None
    )
    tables_by_sheet = defaultdict(list)
    rows_by_ref = defaultdict(list)
    if parsed.structure is not None:
        for sheet in parsed.structure.sheets:
            for block in sheet.blocks:
                if block.logical_table:
                    tables_by_sheet[sheet.name].append(block.logical_table)
                    for row in block.logical_table.rows:
                        rows_by_ref[row.source_ref.key].append(row.role)
    results = []
    try:
        for check in checks:
            kind = check["kind"]
            if kind == "pdf_text":
                text = "\n".join(
                    p.markdown_content for p in parsed.pages if p.page_number == check["page"]
                )
                found = "".join(check["text"].split()) in "".join(text.split())
                results.append(
                    {
                        "check": check,
                        "expected": check["text"],
                        "actual": "found" if found else "not_found",
                        "status": "passed" if found else "failed",
                    }
                )
                continue
            if kind == "pdf_table":
                actual = sum(len(p.tables) for p in parsed.pages if p.page_number == check["page"])
                results.append(
                    {
                        "check": check,
                        "expected": check["minimum"],
                        "actual": actual,
                        "status": "passed" if actual >= check["minimum"] else "failed",
                    }
                )
                continue
            if kind == "drawing_count":
                actual = [
                    len(s.objects)
                    for s in parsed.structure.snapshot.sheets
                    if s.name == check["sheet"]
                ]
                results.append(_semantic_result(check, [check["expected"]], actual))
                continue
            tables = tables_by_sheet[check["sheet"]]
            if kind in {"header", "header_path"}:
                if kind == "header":
                    values = [book[check["sheet"]][c].value for c in check["source_cells"]]
                    if any(v is None or not str(v).strip() for v in values):
                        raise ValueError("Header truth references an empty source cell")
                    expected = [str(v) for v in values]
                else:
                    expected = check["expected"]
                    tables = [
                        t
                        for t in tables
                        if any(
                            range_boundaries(ref.range)[1]
                            <= check["header_row"]
                            <= range_boundaries(ref.range)[3]
                            for ref in t.source_refs
                        )
                    ]
                actual = [
                    c.path for t in tables for c in t.columns if c.coordinate == check["column"]
                ]
                result = _semantic_result(check, [expected], actual)
                # Preserve the original single-header report shape.
                result["expected"] = expected
                results.append(result)
            else:
                actual = rows_by_ref[f"{check['sheet']}!{check['source_range']}"]
                result = _semantic_result(check, [check["expected"]], actual)
                result["expected"] = check["expected"]
                results.append(result)
    finally:
        if book is not None:
            book.close()
    return results


def check_business_records(checks, chunks):
    expected = {
        f"{c['sheet']}!{c['source_range']}"
        for c in checks
        if c["kind"] == "row_role" and c["expected"] in {"data", "total"}
    }
    actual = Counter(
        ref
        for chunk in chunks
        for record in chunk.structured_payload.get("records", [])
        for ref in record.get("source_refs", [])
    )
    missing = sorted(expected - actual.keys())
    duplicates = sorted(ref for ref in expected if actual[ref] > 1)
    return {
        "expected_rows": len(expected),
        "present_once": sum(actual[ref] == 1 for ref in expected),
        "missing": missing,
        "duplicates": duplicates,
        "status": "failed" if missing or duplicates else "passed",
    }


def check_pdf(source_pages, parsed):
    differences = []
    expected_numbers = [p["page"] for p in source_pages]
    if [p.page_number for p in parsed.pages] != expected_numbers:
        differences.append({"code": "page_coverage_mismatch"})
    actual = {p.page_number: p for p in parsed.pages}
    for source in source_pages:
        page = actual.get(source["page"])
        if page is None:
            continue
        if not page.markdown_content.strip():
            differences.append({"code": "empty_page", "page": source["page"]})
        if source["requires_ocr"] and not page.metadata.get("ocr_applied", False):
            differences.append({"code": "scan_without_ocr", "page": source["page"]})
    return {
        "status": "failed" if differences else "passed",
        "source_pages": len(source_pages),
        "parsed_pages": len(parsed.pages),
        "differences": differences,
    }


def _pdf_source(path, output):
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            text_chars = len((page.extract_text() or "").strip())
            area = page.width * page.height
            image_area = max(
                (abs((i["x1"] - i["x0"]) * (i["bottom"] - i["top"])) for i in page.images),
                default=0,
            )
            pages.append(
                {
                    "page": number,
                    "native_text_chars": text_chars,
                    "requires_ocr": image_area / area >= 0.5 and text_chars < 500,
                }
            )
            page.to_image(resolution=110).original.save(output / f"page-{number:03}.png")
    return pages


def _text(value):
    return html.escape("" if value is None else str(value))


def _semantic_review(checks):
    failures = [c for c in checks if c["status"] == "failed"]
    rows = []
    for item in failures:
        check = item["check"]
        location = (
            f"{check['sheet']}: {check.get('column', check.get('source_range', ''))}"
            if "sheet" in check
            else check.get("location", f"PDF 第 {check.get('page', '')} 页")
        )
        rows.append(
            "<tr>"
            + "".join(
                f"<td>{_text(value)}</td>"
                for value in (
                    location,
                    check["kind"],
                    json.dumps(item["expected"], ensure_ascii=False),
                    json.dumps(item["actual"], ensure_ascii=False),
                )
            )
            + "</tr>"
        )
    return (
        f"<h2>已标注语义检查：{len(checks) - len(failures)} 通过，{len(failures)} 失败</h2>"
        "<p>下表列出源文件标注与当前输出的差异。无标注区域仍需复核。</p>"
        '<div class="scroll"><table><tr><th>位置</th><th>检查</th><th>源文件预期</th><th>实际结果</th></tr>'
        + "".join(rows)
        + "</table></div><details><summary>全部标注与检查结果</summary><pre>"
        + _text(json.dumps(checks, ensure_ascii=False, indent=2))
        + "</pre></details>"
    )


def _html_page(title, body):
    return f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>{_text(title)}</title>
<style>body{{font:15px/1.65 system-ui,sans-serif;max-width:1440px;margin:30px auto;padding:0 24px;color:#243143}}
h1,h2{{line-height:1.3}}a{{color:#185eab}}table{{border-collapse:collapse;width:100%;font-size:13px}}
th,td{{border:1px solid #d8e0e8;padding:7px;text-align:left;vertical-align:top;white-space:pre-wrap}}
th{{background:#eff4f8}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f7f9;padding:16px}}
.pair{{display:grid;grid-template-columns:1fr 1fr;gap:24px}}img{{max-width:100%;border:1px solid #ddd}}
details{{margin:20px 0}}.scroll{{overflow:auto}}@media(max-width:800px){{.pair{{display:block}}}}</style>
<h1>{_text(title)}</h1>{body}</html>"""


def _workbook_review(path, parsed):
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=False)
    sections = []
    try:
        for index, sheet in enumerate(book, 1):
            ir = next((s for s in parsed.structure.sheets if s.name == sheet.title), None)
            blocks = ir.blocks if ir else []
            tables = [b.logical_table for b in blocks if b.logical_table]
            roles = defaultdict(list)
            for table in tables:
                for row in table.rows:
                    roles[row.metadata["row_number"]].append(f"{row.source_ref.range}: {row.role}")
            headers = [{"column": c.coordinate, "path": c.path} for t in tables for c in t.columns]
            rows = []
            preview_count = 0
            for row in sheet:
                if not any(c.value is not None for c in row):
                    continue
                values = "<br>".join(
                    f"<b>{c.coordinate}</b> {_text(c.value)}" for c in row if c.value is not None
                )
                row_number = row[0].row
                row_roles = roles[row_number]
                rows.append(
                    f"<tr><td>{row_number}</td><td>{values}</td><td>{_text('; '.join(row_roles))}</td></tr>"
                )
                preview_count += 1
                if preview_count >= 200:
                    break
            sections.append(
                f'<details id="sheet-{index}"><summary>Sheet {index}: {_text(sheet.title)}</summary>'
                "<p>源文件预览最多展示前 200 个非空行；完整解析与差异见 JSON 产物。</p>"
                f"<h3>当前表头（待对照源文件复核）</h3><pre>{_text(json.dumps(headers, ensure_ascii=False, indent=2))}</pre>"
                '<div class="scroll"><table><tr><th>源行</th><th>源文件单元格</th><th>当前解析行角色</th></tr>'
                + "".join(rows)
                + "</table></div></details>"
            )
    finally:
        book.close()
    return "".join(sections)


def _provenance():
    root = Path(__file__).resolve().parents[1]

    def git(*args):
        result = subprocess.run(["git", "-C", str(root), *args], capture_output=True, text=True)
        return result.stdout.strip() if result.returncode == 0 else None

    dependencies = {}
    for name in ("openpyxl", "pdfplumber", "rapidocr-onnxruntime", "onnxruntime", "pypdfium2"):
        try:
            dependencies[name] = version(name)
        except PackageNotFoundError:
            dependencies[name] = None
    return {
        "parser_version": __version__,
        "parser_module": str(Path(langparse.__file__).resolve()),
        "python_executable": sys.executable,
        "runner_sha256": _digest(Path(__file__)),
        "git_head": git("rev-parse", "HEAD"),
        "git_status": git("status", "--porcelain"),
        "dependencies": dependencies,
    }


def run(manifest_path, output_dir, *, pdf_engine="simple", compress=False):
    if pdf_engine not in {"simple", "deepdoc"}:
        raise ValueError("Choose a local PDF engine: simple or deepdoc")
    manifest_path, output_dir = Path(manifest_path).resolve(), Path(output_dir).resolve()
    manifest = _load_manifest(manifest_path)
    sources = [(manifest_path.parent / s["path"]).resolve() for s in manifest["samples"]]
    for source in sources:
        if (
            source == output_dir
            or output_dir in source.parents
            or source.parent in output_dir.parents
        ):
            raise ValueError("Output directory must not overlap input directories")
    if manifest_path == output_dir or output_dir in manifest_path.parents:
        raise ValueError("Output directory must not overlap the manifest")
    output_dir.mkdir(parents=True, exist_ok=False)
    service, results = ParseService(), []
    for sample, source in zip(manifest["samples"], sources, strict=True):
        result = {
            "id": sample["id"],
            "source": str(source),
            "sha256": sample["sha256"],
            "status": "failed",
        }
        results.append(result)
        if not source.is_file():
            result["failure"] = "source_missing"
            continue
        if _digest(source) != sample["sha256"]:
            result["failure"] = "source_hash_mismatch"
            continue
        before = (source.stat().st_size, source.stat().st_mtime_ns, _digest(source))
        output = output_dir / sample["id"]
        output.mkdir()
        started = time.perf_counter()
        try:
            if source.suffix.lower() == ".pdf":
                pages = _pdf_source(source, output)
                result["source_pages"] = pages
                engine_options = (
                    {"model_policy": "require_existing"}
                    if pdf_engine == "deepdoc"
                    else {"enable_ocr": True}
                )
                parsed = service.parse_result(
                    source, engine_name=pdf_engine, chunk=True, **engine_options
                )
                result["engine"] = parsed.engine
                result["source_comparison"] = check_pdf(pages, parsed)
                body = "".join(
                    f'<h2 id="page-{p.page_number}">第 {p.page_number} 页</h2><div class="pair">'
                    f'<img src="page-{p.page_number:03}.png" alt="源 PDF 页面">'
                    f"<pre>{_text(p.markdown_content)}</pre></div>"
                    for p in parsed.pages
                )
            elif source.suffix.lower() == ".csv":
                parsed = service.parse_result(source)
                result["source_comparison"] = compare_csv(
                    source, parsed, sample["semantic_checks"][0]
                )
                body = (
                    "<h2>分隔文本解析预览</h2><pre>"
                    + _text(parsed.markdown_content[:20000])
                    + "</pre>"
                )
            else:
                parsed = service.parse_result(
                    source,
                    chunk=False,
                    chunk_profile="retrieval",
                    workbook_disambiguation=WorkbookDisambiguation.off(),
                )
                result["source_comparison"] = compare_workbook(source, parsed)
                parsed.chunks = service.chunk_result(parsed, chunk_profile="retrieval")
                analysis = service.chunk_result(parsed, chunk_profile="analysis")
                result["business_records"] = check_business_records(
                    sample["semantic_checks"], analysis
                )
                result["chunks"] = {
                    "retrieval": check_chunks(parsed, parsed.chunks),
                    "analysis": check_chunks(parsed, analysis),
                }
                result["block_counts"] = dict(
                    Counter(b.kind for s in parsed.structure.sheets for b in s.blocks)
                )
                result["diagnostics"] = asdict(parsed.diagnostics)
                result["no_model_calls"] = not parsed.diagnostics.model_calls
                body = _workbook_review(source, parsed)
                _write_json(
                    output / ("analysis-chunks.json.gz" if compress else "analysis-chunks.json"),
                    analysis,
                )
                del analysis
            semantic = (
                result["source_comparison"]["semantic_checks"]
                if source.suffix.lower() == ".csv"
                else check_semantics(source, parsed, sample["semantic_checks"])
            )
            if "business_records" in result and sample["semantic_checks"]:
                records = result["business_records"]
                semantic.extend(
                    [
                        _semantic_result(
                            {"kind": "business_records_missing", "location": source.name},
                            [],
                            records["missing"],
                        ),
                        _semantic_result(
                            {"kind": "business_records_duplicate", "location": source.name},
                            [],
                            records["duplicates"],
                        ),
                        _semantic_result(
                            {"kind": "no_model_calls", "location": source.name},
                            True,
                            result["no_model_calls"],
                        ),
                    ]
                )
            result["semantic_checks"] = semantic
            result["semantic_scope"] = "partial_source_annotations" if semantic else "unlabelled"
            failed = result["source_comparison"]["status"] != "passed" or any(
                c["status"] == "failed" for c in semantic
            )
            if "chunks" in result:
                failed |= any(c["status"] != "passed" for c in result["chunks"].values())
                failed |= (
                    not result["no_model_calls"]
                    or parsed.diagnostics.status != "success"
                    or bool(parsed.diagnostics.errors)
                )
            result["status"] = "failed" if failed else "passed"
            _write_json(output / ("parsed.json.gz" if compress else "parsed.json"), parsed)
            (output / "parsed.md").write_text(parsed.markdown_content, encoding="utf-8")
            heading = "<p>仅限本地复核。源文件由独立读取路径展示；解析输出不作为正确答案。</p>"
            heading += _semantic_review(semantic)
            (output / "review.html").write_text(
                _html_page(sample["id"], heading + body), encoding="utf-8"
            )
            result["review"] = f"{sample['id']}/review.html"
        except Exception as exc:
            result["failure"] = type(exc).__name__
            result["status"] = "failed"
        finally:
            result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
            result["source_unchanged"] = source.is_file() and before == (
                source.stat().st_size,
                source.stat().st_mtime_ns,
                _digest(source),
            )
            if not result["source_unchanged"]:
                result.update(status="failed", failure="source_changed_during_run")
        _write_json(output / "checks.json", result)
        print(f"{sample['id']}: {result['status']} ({result['elapsed_seconds']}s)", flush=True)
    summary = {
        "schema_version": 1,
        "split": "regression",
        "pdf_engine": pdf_engine,
        "manifest_sha256": _digest(manifest_path),
        "status": "passed" if all(r["status"] == "passed" for r in results) else "failed",
        "business_review_status": "pending",
        "production_ready": False,
        "semantic_check_count": sum(len(r.get("semantic_checks", [])) for r in results),
        "semantic_failure_count": sum(
            c["status"] == "failed" for r in results for c in r.get("semantic_checks", [])
        ),
        "provenance": _provenance(),
        "samples": results,
    }
    _write_json(output_dir / "summary.json", summary)
    links = "".join(
        f"<li>{_text(Path(r['source']).name)} ({_text(r['id'])}): {_text(r['status'])} "
        + (
            f'<a href="{r["review"]}">源文件与解析结果对照</a>'
            if r.get("review")
            else _text(r.get("failure"))
        )
        + "</li>"
        for r in results
    )
    body = f"<p>自动检查：{summary['status']}。业务复核：pending。生产认证：未通过。</p>"
    body += f"<p>已标注语义检查：{summary['semantic_check_count']} 项，其中 {summary['semantic_failure_count']} 项失败。每项差异均保留预期与实际结果。</p>"
    body += "<p>这是已知真实样本的回归报告。事实一致不代表业务结构正确；未标注区域、OCR 全文准确率和生产场景覆盖仍未认证。</p>"
    (output_dir / "index.html").write_text(
        _html_page("真实样本回归报告", body + "<ul>" + links + "</ul>"), encoding="utf-8"
    )
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    init = commands.add_parser(
        "init", help="Pin local source hashes without generating semantic truth"
    )
    init.add_argument("--manifest", type=Path, required=True)
    init.add_argument("paths", type=Path, nargs="+")
    execute = commands.add_parser("run", help="Run all pinned sources and retain review artifacts")
    execute.add_argument("--manifest", type=Path, required=True)
    execute.add_argument("--output-dir", type=Path, required=True)
    execute.add_argument(
        "--compress", action="store_true", help="Gzip full parsed/chunk JSON artifacts"
    )
    execute.add_argument("--pdf-engine", choices=("simple", "deepdoc"), default="simple")
    args = parser.parse_args(argv)
    try:
        if args.command == "init":
            create_manifest(args.paths, args.manifest)
            return 0
        result = run(
            args.manifest, args.output_dir, pdf_engine=args.pdf_engine, compress=args.compress
        )
        print(f"{result['status']}: {args.output_dir / 'index.html'}")
        return 0 if result["status"] == "passed" else 1
    except (ValueError, OSError) as exc:
        parser.exit(2, f"Regression setup failed ({type(exc).__name__}): {exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
