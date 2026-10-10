import pytest

from langparse.chunkers import FixedTokenChunker, SlidingWindowChunker
from langparse.chunkers.registry import create_chunker, register_chunker
from langparse.types import Document


def test_fixed_lexical_tokens_round_trip_and_budget():
    text = "# 中文\nhello world!  👋\n"
    chunks = FixedTokenChunker(max_chunk_size=3).chunk(Document(content=text))
    assert "".join(c.content for c in chunks) == text
    assert all(c.metadata["token_count"] <= 3 for c in chunks)
    assert all(c.metadata["tokenizer"] == "lexical" for c in chunks)


def test_injected_codec_and_overlap():
    chunks = FixedTokenChunker(max_chunk_size=3, overlap=1, encoder=list, decoder="".join).chunk(
        Document(content="abcdef")
    )
    assert [c.content for c in chunks] == ["abc", "cde", "ef"]


def test_sliding_window_crosses_blocks_and_preserves_metadata():
    doc = Document(content="ab\n\ncd", metadata={"source": "demo"})
    chunks = SlidingWindowChunker(max_chunk_size=4, overlap=2).chunk(doc)
    assert [c.content for c in chunks] == ["ab\n\n", "\n\ncd"]
    assert chunks[1].metadata["char_start"] == 2
    assert chunks[0].metadata["source"] == "demo"
    assert doc.metadata == {"source": "demo"}


@pytest.mark.parametrize("cls", [FixedTokenChunker, SlidingWindowChunker])
@pytest.mark.parametrize(
    "options", [{"max_chunk_size": 0}, {"overlap": -1}, {"max_chunk_size": 2, "overlap": 2}]
)
def test_invalid_budget(cls, options):
    with pytest.raises(ValueError):
        cls(**options)


def test_codec_pair_required():
    with pytest.raises(ValueError, match="encoder and decoder"):
        FixedTokenChunker(encoder=list)


def test_registry_rejects_duplicate_and_unknown():
    assert isinstance(create_chunker("sliding-window"), SlidingWindowChunker)
    with pytest.raises(ValueError, match="already registered"):
        register_chunker("semantic", SlidingWindowChunker)
    with pytest.raises(ValueError, match="Unknown chunk strategy"):
        create_chunker("missing")


def test_service_and_cli_strategy(tmp_path, capsys):
    import json

    from langparse.cli import main
    from langparse.services.parse_service import ParseService

    path = tmp_path / "sample.md"
    path.write_text("abcdef", encoding="utf-8")
    parsed = ParseService().parse_result(
        path,
        chunk=True,
        chunk_strategy="sliding-window",
        chunk_options={"max_chunk_size": 3, "overlap": 1},
    )
    assert [c.content for c in parsed.chunks] == ["abc", "cde", "ef"]
    assert (
        main(
            [
                "parse",
                str(path),
                "--chunk",
                "--chunk-strategy",
                "sliding-window",
                "--chunk-size",
                "3",
                "--chunk-overlap",
                "1",
                "--format",
                "json",
            ]
        )
        == 0
    )
    payload = json.loads(capsys.readouterr().out)
    assert [c["content"] for c in payload["chunks"]] == ["abc", "cde", "ef"]


def test_workbook_strategy_keeps_structural_evidence(tmp_path):
    from openpyxl import Workbook

    from langparse.services.parse_service import ParseService

    path = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.append(["Product", "Price"])
    book.active.append(["A", 10])
    book.save(path)
    service = ParseService()
    parsed = service.parse_result(path)
    expected = service.chunk_result(parsed)
    assert expected
    actual = service.chunk_result(
        parsed, chunk_strategy="sliding-window", chunk_options={"max_chunk_size": 1}
    )
    assert [(c.content, c.metadata) for c in actual] == [(c.content, c.metadata) for c in expected]
    with pytest.raises(ValueError, match="Unknown chunk strategy"):
        service.chunk_result(parsed, chunk_strategy="missing")


def test_batch_strategy_does_not_leak_to_engine(tmp_path, monkeypatch):
    from langparse.services.batch_service import BatchParseService
    from langparse.services.parse_service import ParseService

    path = tmp_path / "sample.md"
    path.write_text("abcdef", encoding="utf-8")
    seen = []
    monkeypatch.setattr(
        ParseService, "create_engine", lambda self, name, **kw: seen.append(kw) or object()
    )
    result = BatchParseService().run(
        [path],
        output_dir=None,
        chunk=True,
        chunk_strategy="fixed-token",
        chunk_options={"max_chunk_size": 2},
    )
    assert seen == [{}]
    assert result.items[0].status == "success"
