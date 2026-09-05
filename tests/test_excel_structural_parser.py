from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from langparse.parsers.excel_parser import ExcelParser
from langparse.services.parse_service import ParseService
from langparse.workbooks.types import WorkbookIR


def test_xlsx_parser_returns_structure_without_fake_pages(sample_excel_file):
    parser = ExcelParser()

    parsed = parser.parse_result(sample_excel_file)
    rendered = parser.parse(sample_excel_file)

    assert isinstance(parsed.structure, WorkbookIR)
    assert parsed.paginated is False
    assert parsed.diagnostics is not None
    assert parsed.diagnostics.coverage_ratio == 1.0
    assert "<!-- page_number:" not in rendered.content
    assert "### Sheet: Sheet1" not in rendered.content
    assert "## Sheet: Sheet1" in rendered.content
    assert parsed.markdown_content == "\n\n".join(page.markdown_content for page in parsed.pages)


def test_xlsx_parser_does_not_emit_unnamed_headers(tmp_path):
    path = tmp_path / "cover.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Label"
    sheet["B1"] = "Title"
    workbook.save(path)

    parsed = ExcelParser().parse_result(path)

    assert "Unnamed:" not in parsed.markdown_content
    assert parsed.pages[0].tables[0]["rows"][0] == ["A", "B"]
    assert parsed.pages[0].tables[0]["rows"][1] == ["Label", "Title"]


def test_native_anchors_flow_through_blocks_markdown_and_chunks(tmp_path):
    path = tmp_path / "adjacent.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in (
        ["Name", "Value", "Item", "Amount"],
        ["Alpha", 1, "First", 10],
        ["Beta", 2, "Second", 20],
    ):
        sheet.append(row)
    sheet.add_table(Table(displayName="LeftTable", ref="A1:B3"))
    sheet.add_table(Table(displayName="RightTable", ref="C1:D3"))
    workbook.save(path)

    parsed = ExcelParser().parse_result(path)
    chunks = ParseService().chunk_result(parsed)
    blocks = parsed.structure.sheets[0].blocks

    assert [block.source_refs[0].range for block in blocks] == ["A1:B3", "C1:D3"]
    assert all(block.metadata["region_reason_codes"] == ["native_table_anchor"] for block in blocks)
    assert "| Name | Value |" in parsed.markdown_content
    assert "| Item | Amount |" in parsed.markdown_content
    assert [chunk.metadata["fragment_ranges"] for chunk in chunks] == [
        ["Data!A1:B3"],
        ["Data!C1:D3"],
    ]
    assert parsed.diagnostics is not None
    assert parsed.diagnostics.region_diagnostics == [
        {
            "sheet_name": "Data",
            "range": "A1:B3",
            "reason_codes": ["native_table_anchor"],
            "confidence": 0.98,
            "conflicts": [],
        },
        {
            "sheet_name": "Data",
            "range": "C1:D3",
            "reason_codes": ["native_table_anchor"],
            "confidence": 0.98,
            "conflicts": [],
        },
    ]
