import json

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from langparse.parsers.excel_parser import ExcelParser
from langparse.services.parse_service import ParseService


def _parsed(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Amount", "Result"])
    sheet.append([10, "=A2*2"])
    sheet.append([20, "=A3*2"])
    sheet.add_table(Table(displayName="Records", ref="A1:B3"))
    path = tmp_path / "bundle.xlsx"
    workbook.save(path)
    return ExcelParser().parse_result(path)


def test_bundle_queries_and_explicit_omission_policy(tmp_path):
    from langparse.workbooks.bundle import WorkbookBundle

    parsed = _parsed(tmp_path)
    bundle = WorkbookBundle.from_result(parsed, max_rows=1)
    assert bundle.tables()[0]["row_count"] == 3
    assert bundle.tables()[0]["omitted_rows"] == 2
    assert bundle.cells(sheet="Data", source_range="B2")[0]["formula"] == "=A2*2"
    assert bundle.relations(source_ref="Data!B2")[0]["target"] == "Data!A2"
    payload = bundle.to_dict()
    assert payload["schema_version"] == 1
    assert "snapshot" not in payload
    assert payload["facts"]["sheets"][0]["cell_count"] == 6
    assert payload["facts"]["sheets"][0]["cells"] is None
    assert payload["provenance"]["omissions"]["cells"] is True
    loaded = WorkbookBundle.from_json(bundle.to_json())
    assert loaded.blocks(kind="logical_table")[0]["source_refs"] == ["Data!A1:B3"]
    with pytest.raises(ValueError, match="omitted"):
        loaded.cells(sheet="Data")


def test_bundle_round_trip_unknown_fields_and_schema_validation(tmp_path):
    from langparse.workbooks.bundle import WorkbookBundle

    bundle = WorkbookBundle.from_result(_parsed(tmp_path), include_cells=True)
    payload = bundle.to_dict()
    payload["future_extension"] = {"enabled": True}
    loaded = WorkbookBundle.from_json(json.dumps(dict(reversed(list(payload.items())))))
    assert loaded.to_dict()["future_extension"] == {"enabled": True}
    assert loaded.cells(sheet="Data", source_range="A2")[0]["raw_value"] == 10
    payload["schema_version"] = 999
    with pytest.raises(ValueError, match="schema"):
        WorkbookBundle.from_json(json.dumps(payload))
    for invalid in (0, -1, True):
        with pytest.raises(ValueError):
            WorkbookBundle.from_result(_parsed(tmp_path), max_rows=invalid)


def test_bundle_json_output_matches_library_export(tmp_path):
    from langparse.workbooks.bundle import WorkbookBundle

    parsed = _parsed(tmp_path)
    output = ParseService().render_output(parsed, "workbook-json")
    assert json.loads(output) == WorkbookBundle.from_result(parsed).to_dict()


def test_bundle_cli_matches_api_and_rejects_invalid_source(tmp_path, capsys):
    from langparse.cli import main
    from langparse.workbooks.bundle import WorkbookBundle

    parsed = _parsed(tmp_path)
    assert main(["parse", parsed.source, "--format", "workbook-json"]) == 0
    assert json.loads(capsys.readouterr().out) == WorkbookBundle.from_result(parsed).to_dict()
    payload = WorkbookBundle.from_result(parsed).to_dict()
    payload["structures"]["blocks"][0]["source_refs"] = ["Missing!A1"]
    with pytest.raises(ValueError, match="source"):
        WorkbookBundle.from_json(json.dumps(payload))
