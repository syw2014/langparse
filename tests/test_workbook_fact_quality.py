from copy import deepcopy

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from langparse.parsers.excel_parser import ExcelParser
from langparse.workbooks.quality.evaluator import evaluate_workbook_result
from langparse.workbooks.quality.schema import WorkbookQualityManifestError, _load_expectation


def _fixture(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Amount", "Result"])
    sheet.append([10, "=Rate"])
    sheet.add_table(Table(displayName="Records", ref="A1:B2"))
    hidden = book.create_sheet("Inputs")
    hidden["A1"] = 2
    hidden.sheet_state = "hidden"
    book.defined_names.add(DefinedName("Rate", attr_text="'Inputs'!$A$1"))
    path = tmp_path / "facts.xlsx"
    book.save(path)
    truth = {
        "sheets": [{"name": "Data", "blocks": []}, {"name": "Inputs", "blocks": []}],
        "continuations": [],
        "required_source_refs": [],
        "objects": [],
        "facts": {
            "formulas": [{"source_ref": "Data!B2", "formula": "=Rate"}],
            "defined_names": [{"name": "Rate", "definition": "'Inputs'!$A$1", "scope_sheet": None}],
            "tables": [
                {"name": "Records", "source_ref": "Data!A1:B2", "columns": ["Amount", "Result"]}
            ],
            "sheet_visibility": [
                {"sheet": "Data", "state": "visible"},
                {"sheet": "Inputs", "state": "hidden"},
            ],
            "dependencies": [
                {
                    "dependent": "Data!B2",
                    "target": "Inputs!A1",
                    "status": "resolved",
                    "reference": "Rate",
                }
            ],
        },
    }
    return truth, ExcelParser().parse_result(path)


def test_fact_truth_scores_and_detects_missing_formula_dependency(tmp_path):
    truth, parsed = _fixture(tmp_path)
    expected = _load_expectation(truth)
    scores = evaluate_workbook_result(expected, parsed)
    assert (
        scores.formula_accuracy == scores.named_range_accuracy == scores.excel_table_accuracy == 1.0
    )
    assert (
        scores.visibility_accuracy == scores.dependency_precision == scores.dependency_recall == 1.0
    )
    parsed.structure.snapshot.reference_facts.formulas.clear()
    parsed.structure.lineage.edges.clear()
    scores = evaluate_workbook_result(expected, parsed)
    assert scores.formula_accuracy == 0.0
    assert scores.dependency_recall == 0.0


def test_missing_fact_truth_is_unscored(tmp_path):
    truth, parsed = _fixture(tmp_path)
    del truth["facts"]
    scores = evaluate_workbook_result(_load_expectation(truth), parsed)
    assert scores.formula_accuracy is None
    assert scores.visibility_accuracy is None
    assert scores.dependency_recall is None


def test_fact_truth_rejects_unknown_sheet_duplicate_and_bad_visibility(tmp_path):
    truth, _ = _fixture(tmp_path)
    for key, replacement in [
        ("formulas", [{"source_ref": "Missing!A1", "formula": "=1"}]),
        ("sheet_visibility", [{"sheet": "Inputs", "state": "invisible"}]),
        ("dependencies", truth["facts"]["dependencies"] * 2),
    ]:
        broken = deepcopy(truth)
        broken["facts"][key] = replacement
        with pytest.raises(WorkbookQualityManifestError):
            _load_expectation(broken)


def test_fact_metric_participates_in_non_regression_gate(tmp_path):
    import hashlib
    import json

    from langparse.services.workbook_quality_benchmark import WorkbookQualityBenchmarkService

    truth, _ = _fixture(tmp_path)
    truth["facts"]["formulas"][0]["formula"] = "=Rate+1"
    manifest = {
        "schema_version": 1,
        "dataset_id": "fact-truth",
        "dataset_version": "1",
        "split": "tuning",
        "source_root": ".",
        "quality_gate": {"minimum": {"formula_accuracy": 1.0}, "maximum": {}},
        "samples": [
            {
                "sample_id": "fact",
                "path": "facts.xlsx",
                "sha256": "sha256:"
                + hashlib.sha256((tmp_path / "facts.xlsx").read_bytes()).hexdigest(),
                "expectation": truth,
            }
        ],
    }
    path = tmp_path / "manifest.json"
    path.write_text(json.dumps(manifest))
    report = WorkbookQualityBenchmarkService().run(
        path, output_dir=tmp_path.parent / (tmp_path.name + "-report")
    )
    assert report.summary["status"] == "failed"
    assert report.summary["gate_reasons"] == ["formula_accuracy_below_minimum"]
