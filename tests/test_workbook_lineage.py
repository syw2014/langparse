from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from langparse.workbooks.adapters import OOXMLWorkbookAdapter
from langparse.workbooks.lineage import build_lineage
from langparse.workbooks.types import SourceRef


def test_reference_facts_and_lineage(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Amount", "Result"])
    sheet.append([10, "=SUM(Sales[Amount])+Rate"])
    sheet.append([20, "=B2"])
    sheet.add_table(Table(displayName="Sales", ref="A1:A3"))
    other = book.create_sheet("Other Sheet")
    other["A1"] = "=Data!B3"
    other["B1"] = '=INDIRECT("Data!A2")'
    other["C1"] = "='[outside.xlsx]Data'!A1"
    other["D1"] = "=MissingName"
    other["E1"] = "=F1"
    other["F1"] = "=E1"
    book.defined_names.add(DefinedName("Rate", attr_text="'Data'!$A$2"))
    path = tmp_path / "refs.xlsx"
    book.save(path)
    snapshot = OOXMLWorkbookAdapter().snapshot(path)
    facts = snapshot.reference_facts
    assert facts.defined_names[0].definition == "'Data'!$A$2"
    assert facts.tables[0].columns == ["Amount"]
    assert facts.external_references[0].workbook == "outside.xlsx"
    lineage = build_lineage(snapshot)
    deps = lineage.dependencies_of(SourceRef("Data", "B2"))
    assert {edge.target.key for edge in deps} == {"Data!A2:A3", "Data!A2"}
    assert all(edge.source_refs and edge.status == "resolved" for edge in deps)
    assert lineage.dependents_of(SourceRef("Data", "A2"))
    assert lineage.dependencies_of(SourceRef("Other Sheet", "A1"))[0].target.key == "Data!B3"
    codes = {item.code for item in lineage.diagnostics}
    assert {"dynamic_reference", "unresolved_reference", "circular_reference"} <= codes


def test_local_names_structured_rows_and_explicit_degradation(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Amount", "Result"])
    sheet.append([10, "=Sales[@Amount]"])
    sheet.append([20, "=Rate"])
    sheet.add_table(Table(displayName="Sales", ref="A1:B3"))
    book.defined_names.add(DefinedName("Rate", attr_text="'Data'!$A$2"))
    sheet.defined_names.add(DefinedName("Rate", attr_text="'Data'!$A$3"))
    book.defined_names.add(DefinedName("Moving", attr_text="OFFSET(Data!A2,0,0)"))
    sheet["C1"] = "=SUM(Sales)"
    sheet["C2"] = "=Moving"
    sheet["C3"] = "=#REF!"
    sheet["C4"] = "=SUM(A:A)"
    sheet["C5"] = "=SUM(Sales[#Headers])"
    path = tmp_path / "scopes.xlsx"
    book.save(path)
    snapshot = OOXMLWorkbookAdapter().snapshot(path)
    lineage = build_lineage(snapshot)
    assert lineage.dependencies_of(SourceRef("Data", "B2"))[0].target.key == "Data!A2:A2"
    assert lineage.dependencies_of(SourceRef("Data", "B3"))[0].target.key == "Data!A3"
    assert lineage.dependencies_of(SourceRef("Data", "C1"))[0].target.key == "Data!A2:B3"
    assert lineage.dependencies_of(SourceRef("Data", "C5"))[0].target.key == "Data!A1:B1"
    diagnostics = {(d.code, d.source_refs[0].range) for d in lineage.diagnostics}
    assert ("dynamic_reference", "C2") in diagnostics
    assert ("invalid_reference", "C3") in diagnostics
    assert ("unsupported_reference", "C4") in diagnostics


def test_lineage_reaches_parsed_result_json_and_analysis_chunks(tmp_path):
    import json

    from langparse.parsers.excel_parser import ExcelParser
    from langparse.services.parse_service import ParseService
    from langparse.workbooks.assembly import assemble_baseline

    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Amount", "Result"])
    sheet.append([10, "=A2*2"])
    sheet.append([20, "=A3*2"])
    sheet.add_table(Table(displayName="Records", ref="A1:B3"))
    path = tmp_path / "consume.xlsx"
    book.save(path)
    parsed = ExcelParser().parse_result(path)
    assert (
        parsed.structure.lineage.dependencies_of(SourceRef("Data", "B2"))[0].target.key == "Data!A2"
    )
    baseline, _ = assemble_baseline(parsed.structure.snapshot)
    assert baseline.lineage.dependencies_of(SourceRef("Data", "B3"))[0].target.key == "Data!A3"
    service = ParseService()
    chunks = service.chunk_result(parsed, chunk_profile="analysis")
    dependencies = [edge for chunk in chunks for edge in chunk.structured_payload["dependencies"]]
    assert {edge["dependent"]["range"] for edge in dependencies} == {"B2", "B3"}
    payload = json.loads(service.render_output(parsed, "json", chunks))
    assert payload["structure"]["lineage"]["edges"][0]["status"] == "resolved"
    assert payload["structure"]["snapshot"]["reference_facts"]["tables"][0]["name"] == "Records"


def test_quoted_sheet_apostrophe_is_retained(tmp_path):
    book = Workbook()
    book.active.title = "O'"
    book.active["A1"] = 2
    other = book.create_sheet("'Start")
    other["B1"] = "='O'''!A1"
    book.active["B1"] = "='''Start'!B1"
    path = tmp_path / "quoted.xlsx"
    book.save(path)
    lineage = build_lineage(OOXMLWorkbookAdapter().snapshot(path))
    assert lineage.dependencies_of(SourceRef("'Start", "B1"))[0].target == SourceRef("O'", "A1")
    assert lineage.dependencies_of(SourceRef("O'", "B1"))[0].target == SourceRef("'Start", "B1")


def test_single_cell_lineage_avoids_formula_cross_product(tmp_path, monkeypatch):
    import langparse.workbooks.lineage as module

    book = Workbook()
    for row in range(1, 31):
        book.active.cell(row, 1, row)
        book.active.cell(row, 2, f"=A{row}")
    book.active["C1"] = "=C2"
    book.active["C2"] = "=C1"
    path = tmp_path / "linear.xlsx"
    book.save(path)
    snapshot = OOXMLWorkbookAdapter().snapshot(path)
    calls = []
    original = module.overlaps

    def tracked(first, second):
        calls.append((first, second))
        return original(first, second)

    monkeypatch.setattr(module, "overlaps", tracked)
    lineage = module.build_lineage(snapshot)
    assert len(lineage.edges) == 32
    assert any(d.code == "circular_reference" for d in lineage.diagnostics)
    assert not calls
