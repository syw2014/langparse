from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image
from PIL import Image as PILImage


def test_extract_chart_series_axes_and_image_identity(tmp_path):
    from langparse.workbooks.objects import extract_objects

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Data"
    for row in [("Month", "Sales"), ("Jan", 12), ("Feb", 20)]:
        sheet.append(row)
    chart = BarChart()
    chart.title = "Monthly sales"
    chart.x_axis.title = "Month"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    chart += LineChart()
    sheet.add_chart(chart, "D2")
    path = tmp_path / "image.png"
    PILImage.new("RGB", (8, 6), "white").save(path)
    sheet.add_image(Image(path), "D12")
    objects = extract_objects(sheet)
    assert objects[0].title == "Monthly sales"
    assert objects[0].chart_types == ["BarChart", "LineChart"]
    assert objects[0].series[0].values == "'Data'!$B$2:$B$3"
    assert objects[0].series[0].categories == "'Data'!$A$2:$A$3"
    assert objects[0].axes[0]["title"] == "Month"
    assert objects[1].width == 8 and objects[1].height == 6
    assert objects[1].media_id
    assert objects[1].enhancement_status == "not_requested"


def test_object_only_sheet_is_rendered_and_chunked(tmp_path):
    from langparse.chunkers.workbook import WorkbookStructuralChunker
    from langparse.parsers.excel_parser import ExcelParser

    wb = Workbook()
    sheet = wb.active
    chart = BarChart()
    chart.title = "Empty chart"
    sheet.add_chart(chart, "C4")
    path = tmp_path / "chart.xlsx"
    wb.save(path)
    parsed = ExcelParser().parse_result(path)
    blocks = parsed.structure.sheets[0].blocks
    assert len(blocks) == 1 and blocks[0].kind == "chart"
    assert "Empty chart" in parsed.markdown_content
    chunks = WorkbookStructuralChunker(profile="analysis").chunk(parsed)
    assert chunks[0].metadata["source_ranges"] == ["Sheet!C4"]
    assert chunks[0].structured_payload["object"]["title"] == "Empty chart"
    assert parsed.diagnostics.object_coverage_ratio == 1.0


def test_unresolved_object_is_diagnostic_not_success():
    from langparse.types import ParseDiagnostics
    from langparse.workbooks.objects import attach_objects
    from langparse.workbooks.types import SheetIR, SheetSnapshot, WorkbookIR, WorkbookSnapshot

    snapshot = WorkbookSnapshot(
        "x.xlsx",
        "x.xlsx",
        [
            SheetSnapshot(
                "Data",
                0,
                objects=[
                    {
                        "kind": "image",
                        "anchor": None,
                        "object_id": "bad",
                        "diagnostics": ["unresolved_anchor"],
                    }
                ],
            )
        ],
    )
    ir = WorkbookIR(
        kind="workbook",
        workbook_id="w",
        source="x.xlsx",
        snapshot=snapshot,
        sheets=[SheetIR("s", "Data", 0)],
    )
    diagnostics = ParseDiagnostics()
    attach_objects(snapshot, ir, diagnostics)
    assert diagnostics.status == "partial"
    assert diagnostics.object_coverage_ratio == 0.0
    assert diagnostics.warnings


def test_object_chunk_rejects_unknown_sheet_source(tmp_path):
    import pytest

    from langparse.chunkers.workbook import WorkbookStructuralChunker
    from langparse.parsers.excel_parser import ExcelParser

    workbook = Workbook()
    workbook.active.add_chart(BarChart(), "C4")
    path = tmp_path / "object.xlsx"
    workbook.save(path)
    parsed = ExcelParser().parse_result(path)
    chunker = WorkbookStructuralChunker()
    chunks = chunker.chunk(parsed)
    chunks[0].metadata["source_ranges"] = ["Missing!A1"]
    with pytest.raises(ValueError, match="unknown sheet"):
        chunker._validate_chunks(parsed, chunks)
