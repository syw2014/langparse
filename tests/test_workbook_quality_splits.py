import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook


def _write_workbook(path: Path, value: object) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Value"])
    sheet.append(["Alpha", value])
    workbook.save(path)


def _manifest_payload(
    *,
    dataset_id: str,
    split: str,
    sample_id: str,
    sample_path: str,
    sha256: str,
) -> dict:
    return {
        "schema_version": 1,
        "dataset_id": dataset_id,
        "dataset_version": "1",
        "split": split,
        "source_root": "fixtures",
        "quality_gate": {
            "minimum": {"block_recall": 1.0},
            "maximum": {"fallback_rate": 0.0},
        },
        "samples": [
            {
                "sample_id": sample_id,
                "path": sample_path,
                "sha256": sha256,
                "expectation": {
                    "sheets": [
                        {
                            "name": "Data",
                            "blocks": [
                                {
                                    "source_range": "A1:B2",
                                    "kind": "logical_table",
                                    "headers": [
                                        {"coordinate": "A", "path": ["Name"]},
                                        {"coordinate": "B", "path": ["Value"]},
                                    ],
                                    "rows": [
                                        {"source_range": "A1:B1", "role": "header"},
                                        {"source_range": "A2:B2", "role": "data"},
                                    ],
                                    "form_fields": [],
                                    "matrix_axes": {"rows": [], "columns": []},
                                }
                            ],
                        }
                    ],
                    "continuations": [],
                    "required_source_refs": ["Data!A1:B2"],
                    "objects": [],
                },
            }
        ],
    }


def _write_split_pair(
    root: Path,
    *,
    holdout_value: object = 2,
    holdout_file: str = "holdout.xlsx",
) -> tuple[Path, Path]:
    """Write one tuning and one holdout dataset under separate source roots."""
    pair: dict[str, Path] = {}
    for name, split, value, filename in (
        ("tuning", "tuning", 1, "tuning.xlsx"),
        ("holdout", "holdout", holdout_value, holdout_file),
    ):
        dataset_root = root / name
        fixtures = dataset_root / "fixtures"
        fixtures.mkdir(parents=True)
        workbook_path = fixtures / filename
        _write_workbook(workbook_path, value)
        manifest_path = dataset_root / "manifest.json"
        manifest_path.write_text(
            json.dumps(
                _manifest_payload(
                    dataset_id=f"quality-{name}",
                    split=split,
                    sample_id=f"{name}-table",
                    sample_path=filename,
                    sha256="sha256:" + hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                )
            ),
            encoding="utf-8",
        )
        pair[name] = manifest_path
    return pair["tuning"], pair["holdout"]


def _load(path: Path):
    from langparse.workbooks.quality.schema import load_workbook_quality_manifest

    return load_workbook_quality_manifest(path)


def test_verified_isolation_records_privacy_safe_provenance(tmp_path: Path):
    from langparse.workbooks.quality.splits import verify_split_isolation

    tuning_path, holdout_path = _write_split_pair(tmp_path)
    tuning = _load(tuning_path)
    holdout = _load(holdout_path)

    record = verify_split_isolation(holdout, (tuning,))
    payload = record.as_payload()

    assert payload["holdout_dataset_id"] == "quality-holdout"
    assert payload["holdout_dataset_digest"] == holdout.dataset_digest
    assert payload["holdout_sample_count"] == 1
    assert payload["tuning_datasets"] == [
        {
            "dataset_id": "quality-tuning",
            "dataset_version": "1",
            "dataset_digest": tuning.dataset_digest,
            "sample_count": 1,
        }
    ]
    serialized = json.dumps(payload)
    assert "fixtures" not in serialized
    assert "Alpha" not in serialized


def test_rejects_holdout_sample_repeating_tuning_content(tmp_path: Path):
    from langparse.workbooks.quality.splits import (
        WorkbookSplitIsolationError,
        verify_split_isolation,
    )

    # Same cell values under a different filename: identical content hash.
    tuning_path, holdout_path = _write_split_pair(tmp_path, holdout_value=1)
    tuning = _load(tuning_path)
    holdout = _load(holdout_path)

    with pytest.raises(WorkbookSplitIsolationError, match="repeats tuning content"):
        verify_split_isolation(holdout, (tuning,))


def test_rejects_holdout_sample_reusing_tuning_source_file(tmp_path: Path):
    from langparse.workbooks.quality.splits import (
        WorkbookSplitIsolationError,
        verify_split_isolation,
    )

    tuning_path, _ = _write_split_pair(tmp_path)
    tuning = _load(tuning_path)
    shared = tuning.source_root / "tuning.xlsx"

    # A holdout manifest pointing at the tuning file itself, via its own source root.
    holdout_root = tmp_path / "holdout-alias"
    holdout_root.mkdir()
    (holdout_root / "fixtures").symlink_to(tuning.source_root, target_is_directory=True)
    holdout_path = holdout_root / "manifest.json"
    holdout_path.write_text(
        json.dumps(
            _manifest_payload(
                dataset_id="quality-alias",
                split="holdout",
                sample_id="alias-table",
                sample_path="tuning.xlsx",
                sha256="sha256:" + hashlib.sha256(shared.read_bytes()).hexdigest(),
            )
        ),
        encoding="utf-8",
    )

    with pytest.raises(WorkbookSplitIsolationError):
        verify_split_isolation(_load(holdout_path), (tuning,))


def test_rejects_wrong_split_labels_and_empty_reference_set(tmp_path: Path):
    from langparse.workbooks.quality.splits import (
        WorkbookSplitIsolationError,
        verify_split_isolation,
    )

    tuning_path, holdout_path = _write_split_pair(tmp_path)
    tuning = _load(tuning_path)
    holdout = _load(holdout_path)

    with pytest.raises(WorkbookSplitIsolationError, match="must declare split holdout"):
        verify_split_isolation(tuning, (tuning,))
    with pytest.raises(WorkbookSplitIsolationError, match="must declare split tuning"):
        verify_split_isolation(holdout, (holdout,))
    with pytest.raises(WorkbookSplitIsolationError, match="at least one tuning manifest"):
        verify_split_isolation(holdout, ())


def test_rejects_duplicate_tuning_dataset_reference(tmp_path: Path):
    from langparse.workbooks.quality.splits import (
        WorkbookSplitIsolationError,
        verify_split_isolation,
    )

    tuning_path, holdout_path = _write_split_pair(tmp_path)
    tuning = _load(tuning_path)

    with pytest.raises(WorkbookSplitIsolationError, match="Duplicate tuning dataset"):
        verify_split_isolation(_load(holdout_path), (tuning, tuning))


def test_benchmark_embeds_isolation_provenance_and_changes_digest(tmp_path: Path):
    from langparse.services.workbook_quality_benchmark import WorkbookQualityBenchmarkService

    tuning_path, holdout_path = _write_split_pair(tmp_path)
    service = WorkbookQualityBenchmarkService()
    output_dir = tmp_path / "reports"

    unverified = service.run(holdout_path, output_dir=output_dir)
    verified = service.run(holdout_path, output_dir=output_dir, tuning_manifests=(tuning_path,))

    assert "split_isolation" not in unverified.summary
    assert verified.summary["split_isolation"]["tuning_datasets"][0]["dataset_id"] == (
        "quality-tuning"
    )
    assert verified.run_digest != unverified.run_digest
    summary = json.loads(
        verified.output_path.joinpath("workbook-quality-summary.json").read_text(encoding="utf-8")
    )
    assert summary["split_isolation"] == verified.summary["split_isolation"]
    markdown = verified.output_path.joinpath("workbook-quality-summary.md").read_text(
        encoding="utf-8"
    )
    assert "Split Isolation Verified Against" in markdown
    assert "quality-tuning@1" in markdown


def test_cli_reports_split_leakage_as_execution_error(tmp_path: Path, capsys):
    from langparse.cli import main

    tuning_path, holdout_path = _write_split_pair(tmp_path, holdout_value=1)

    exit_code = main(
        [
            "benchmark-workbook-quality",
            str(holdout_path),
            "--output-dir",
            str(tmp_path / "cli-reports"),
            "--tuning-manifest",
            str(tuning_path),
        ]
    )

    assert exit_code == 2
    assert "repeats tuning content" in capsys.readouterr().err
