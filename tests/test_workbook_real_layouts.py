"""Small synthetic regressions distilled from private acceptance layouts."""

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from langparse.parsers.excel_parser import ExcelParser


def _parse(tmp_path, workbook):
    path = tmp_path / "layout.xlsx"
    workbook.save(path)
    return ExcelParser().parse_result(path)


def _header(sheet, values):
    sheet.append(values)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="336699")


def _table(parsed):
    tables = [b.logical_table for b in parsed.structure.sheets[0].blocks if b.logical_table]
    assert len(tables) == 1
    return tables[0]


def test_blank_input_column_keeps_calculated_amount_with_record(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    _header(sheet, ["序号", "模块", "分组", "内容", None, "金额"])
    for number in (1, 2, 3):
        sheet.append([number, "平台", "业务", "功能", None, f"=E{number + 1}*2"])
    parsed = _parse(tmp_path, workbook)
    table = _table(parsed)

    assert table.source_refs[0].range == "A1:F4"
    assert [c.path for c in table.columns] == [["序号"], ["模块"], ["分组"], ["内容"], [], ["金额"]]
    assert [r.role for r in table.rows] == ["header", "data", "data", "data"]


def test_grouped_quote_inherits_schema_across_blank_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    _header(sheet, ["序号", "大类", "分项", "说明", "占比", "金额", "备注"])
    sheet.append(["一、平台", None, None, "预算", None, 100])
    sheet.merge_cells("A2:C2")
    sheet.append([1, "平台", "开发", "说明", 0.5, 50])
    sheet.append([2, "平台", "维护", "说明", 0.5, 50])
    sheet.append(["平台 小计", None, None, None, 1, "=SUM(F3:F4)"])
    sheet.merge_cells("A5:D5")
    sheet.append([])
    sheet.append(["二、服务", None, None, "预算", None, 100])
    sheet.merge_cells("A7:C7")
    sheet.append([1, "服务", "支持", "说明", 1, 100])
    sheet.append([])
    sheet.append(["项目报价总计", None, None, None, None, "=F5+F8"])
    sheet.merge_cells("A10:D10")
    table = _table(_parse(tmp_path, workbook))

    assert [c.path for c in table.columns] == [
        [v] for v in ["序号", "大类", "分项", "说明", "占比", "金额", "备注"]
    ]
    assert {r.metadata["row_number"]: r.role for r in table.rows if r.role != "unknown"} == {
        1: "header",
        2: "section_header",
        3: "data",
        4: "data",
        5: "total",
        7: "section_header",
        8: "data",
        10: "total",
    }
    assert all(r.source_ref.range.endswith(f"G{r.metadata['row_number']}") for r in table.rows)
    assert table.rows[7].section_path == ["二、服务"]


def test_merged_product_label_does_not_turn_material_lines_into_sections(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    _header(sheet, ["产品", "项目", "单位", "数量", "单价", "金额", "总额"])
    sheet.append(["7-1-202", "主体", "米", 2, 10, 20, 40])
    sheet.append([None, "附件", "米", 1, 10, 10])
    sheet.append([None, None, "米", 1, 10, 10])
    sheet.merge_cells("A2:A4")
    sheet.merge_cells("G2:G4")
    sheet.append(["总计", None, None, None, None, None, 40])
    table = _table(_parse(tmp_path, workbook))

    assert [r.role for r in table.rows] == ["header", "data", "data", "data", "total"]
    assert table.sections == []


def test_nullable_text_records_keep_header_and_do_not_match_total_in_prose(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    _header(sheet, ["标题", "正文", "可选字段", "作者"])
    sheet.append(["一、内容提纲", "本文合计介绍三个例子", None, "张三"])
    sheet.append(["另一条", "描述", "标签", "李四"])
    sheet.append(["第三条", "内容", None, "王五"])
    table = _table(_parse(tmp_path, workbook))

    assert [r.role for r in table.rows] == ["header", "data", "data", "data"]
    assert [c.path for c in table.columns] == [["标题"], ["正文"], ["可选字段"], ["作者"]]
